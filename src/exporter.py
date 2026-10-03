"""Экспорт дайджеста в Markdown и DOCX."""
from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt

from src.models import Digest

logger = logging.getLogger(__name__)


def export_markdown(digest: Digest, path: str | Path) -> Path:
    """Сохраняет дайджест в файл Markdown."""
    path = Path(path)
    path.write_text(digest.to_markdown(), encoding="utf-8")
    logger.info("Markdown сохранён: %s", path)
    return path


def export_docx(digest: Digest, path: str | Path) -> Path:
    """Сохраняет дайджест в файл DOCX."""
    path = Path(path)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading("ИТ/ИБ-дайджест", level=0)

    meta = doc.add_paragraph()
    meta.add_run(f"Интересы: {', '.join(digest.interests)}\n").bold = True
    meta.add_run(f"Дата формирования: {digest.created_at}\n")
    meta.add_run(
        f"Всего собрано: {digest.total_articles}    "
        f"Отобрано: {digest.selected_articles}"
    )

    for cluster in digest.clusters:
        doc.add_heading(f"{cluster.title} ({cluster.size})", level=1)
        for article in cluster.articles:
            doc.add_heading(article.title, level=2)

            meta_line = doc.add_paragraph()
            run = meta_line.add_run(
                f"Источник: {article.source}    "
                f"Релевантность: {article.score:.2f}"
            )
            run.italic = True

            body = article.llm_summary or article.summary
            if body:
                doc.add_paragraph(body.strip())

            link_para = doc.add_paragraph()
            link_run = link_para.add_run(article.link)
            link_run.italic = True

    doc.save(path)
    logger.info("DOCX сохранён: %s", path)
    return path